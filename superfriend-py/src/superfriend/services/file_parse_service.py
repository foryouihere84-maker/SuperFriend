from __future__ import annotations

import structlog
from pathlib import Path
from typing import Any

logger = structlog.get_logger()

MAX_FILE_CONTENT_LENGTH = 500_000
MAX_SINGLE_FILE_LENGTH = 200_000
SMALL_FILE_THRESHOLD = 50 * 1024


class ParseResult:
    def __init__(self, content: str = "", mime_type: str = "", pages: int = 0):
        self.content = content
        self.mime_type = mime_type
        self.pages = pages


class FileParseService:
    async def parse_file(self, file_path: str, mime_type: str | None = None) -> ParseResult:
        path = Path(file_path)
        if not path.exists():
            return ParseResult(content=f"File not found: {file_path}")

        ext = path.suffix.lower()
        if not mime_type:
            mime_type = self._guess_mime(ext)

        try:
            content = await self._parse_by_extension(str(path), ext)
            content = self._truncate(content)
            return ParseResult(content=content, mime_type=mime_type, pages=1)
        except Exception as e:
            logger.error("file_parse_error", file=file_path, error=str(e))
            return ParseResult(content=f"Error parsing file: {e}")

    async def _parse_by_extension(self, file_path: str, ext: str) -> str:
        if ext in (".txt", ".md", ".csv", ".json", ".xml", ".yaml", ".yml", ".log"):
            return await self._parse_text(file_path)

        if ext == ".pdf":
            return await self._parse_pdf(file_path)

        if ext in (".docx", ".doc"):
            return await self._parse_docx(file_path)

        if ext in (".xlsx", ".xls"):
            return await self._parse_xlsx(file_path)

        if ext in (".pptx", ".ppt"):
            return await self._parse_pptx(file_path)

        return await self._parse_text(file_path)

    async def _parse_text(self, file_path: str) -> str:
        encodings = ["utf-8", "gbk", "gb2312", "latin-1"]
        for enc in encodings:
            try:
                with open(file_path, "r", encoding=enc) as f:
                    return f.read()
            except (UnicodeDecodeError, UnicodeError):
                continue
        with open(file_path, "r", encoding="utf-8", errors="replace") as f:
            return f.read()

    async def _parse_pdf(self, file_path: str) -> str:
        try:
            import fitz
            doc = fitz.open(file_path)
            content_parts = []
            for page in doc:
                content_parts.append(page.get_text())
            doc.close()
            return "\n".join(content_parts)
        except ImportError:
            try:
                from PyPDF2 import PdfReader
                reader = PdfReader(file_path)
                return "\n".join(page.extract_text() or "" for page in reader.pages)
            except ImportError:
                logger.warning("pdf_parsers_not_installed")
                return f"[PDF file: {file_path} - install PyMuPDF or PyPDF2 for parsing]"

    async def _parse_docx(self, file_path: str) -> str:
        try:
            import docx
            doc = docx.Document(file_path)
            return "\n".join(p.text for p in doc.paragraphs)
        except ImportError:
            logger.warning("docx_parser_not_installed")
            return f"[DOCX file: {file_path} - install python-docx for parsing]"

    async def _parse_xlsx(self, file_path: str) -> str:
        try:
            import openpyxl
            wb = openpyxl.load_workbook(file_path, data_only=True)
            content_parts = []
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                content_parts.append(f"--- Sheet: {sheet_name} ---")
                for row in ws.iter_rows(values_only=True):
                    content_parts.append("\t".join(str(c) if c is not None else "" for c in row))
                content_parts.append("")
            return "\n".join(content_parts)
        except ImportError:
            logger.warning("xlsx_parser_not_installed")
            return f"[XLSX file: {file_path} - install openpyxl for parsing]"

    async def _parse_pptx(self, file_path: str) -> str:
        try:
            from pptx import Presentation
            prs = Presentation(file_path)
            content_parts = []
            for idx, slide in enumerate(prs.slides, 1):
                content_parts.append(f"--- Slide {idx} ---")
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text:
                        content_parts.append(shape.text)
                content_parts.append("")
            return "\n".join(content_parts)
        except ImportError:
            logger.warning("pptx_parser_not_installed")
            return f"[PPTX file: {file_path} - install python-pptx for parsing]"

    def _truncate(self, content: str) -> str:
        if len(content) > MAX_FILE_CONTENT_LENGTH:
            return content[:MAX_FILE_CONTENT_LENGTH] + "\n... (content truncated)"
        return content

    def _guess_mime(self, ext: str) -> str:
        mime_map = {
            ".pdf": "application/pdf",
            ".doc": "application/msword",
            ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            ".xls": "application/vnd.ms-excel",
            ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            ".ppt": "application/vnd.ms-powerpoint",
            ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            ".txt": "text/plain",
            ".csv": "text/csv",
            ".md": "text/markdown",
            ".json": "application/json",
            ".xml": "application/xml",
            ".jpg": "image/jpeg",
            ".jpeg": "image/jpeg",
            ".png": "image/png",
            ".gif": "image/gif",
            ".webp": "image/webp",
        }
        return mime_map.get(ext, "application/octet-stream")


file_parse_service = FileParseService()